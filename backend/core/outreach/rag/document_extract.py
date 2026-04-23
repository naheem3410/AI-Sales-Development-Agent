"""
document_extract.py
-------------------
Extract plain text from uploaded files for RAG ingestion.
"""

from io import BytesIO
from pathlib import Path


MAX_UPLOAD_BYTES = 15 * 1024 * 1024  # 15 MiB

ALLOWED_EXTENSIONS = frozenset({".pdf", ".docx", ".txt", ".md", ".markdown"})


class DocumentExtractError(Exception):
    """Raised when a file cannot be read or yields no usable text."""


def extract_text_from_upload(filename: str, data: bytes) -> str:
    """
    Return UTF-8 plain text from supported binary/text uploads.

    Supported: .pdf, .docx, .txt, .md, .markdown
    """
    if len(data) > MAX_UPLOAD_BYTES:
        raise DocumentExtractError(
            f"File exceeds maximum size ({MAX_UPLOAD_BYTES // (1024 * 1024)} MiB)."
        )

    ext = Path(filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise DocumentExtractError(
            f"Unsupported file type {ext or '(none)'}. "
            f"Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}"
        )

    if ext in (".txt", ".md", ".markdown"):
        return _decode_plain_text(data)

    if ext == ".pdf":
        return _extract_pdf(data)

    if ext == ".docx":
        return _extract_docx(data)

    raise DocumentExtractError(f"Unhandled extension {ext}")


def _decode_plain_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace").strip()


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise DocumentExtractError(
            "PDF support requires the 'pypdf' package to be installed."
        ) from e

    reader = PdfReader(BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        t = page.extract_text()
        if t:
            parts.append(t.strip())
    text = "\n\n".join(parts).strip()
    if not text:
        raise DocumentExtractError(
            "No text could be extracted from this PDF (it may be scanned images only)."
        )
    return text


def _extract_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as e:
        raise DocumentExtractError(
            "Word support requires the 'python-docx' package to be installed."
        ) from e

    document = docx.Document(BytesIO(data))
    parts: list[str] = []
    for para in document.paragraphs:
        if para.text and para.text.strip():
            parts.append(para.text.strip())
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise DocumentExtractError("No text could be extracted from this Word document.")
    return text


def default_title_from_filename(filename: str) -> str:
    stem = Path(filename or "document").stem.strip() or "document"
    return stem[:200]
